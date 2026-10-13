# -*- coding: utf-8 -*-
"""把官方 Form I docx 原件转成占位符模板 form1_tpl.docx（版式 100% 保留）。
用法：python make_form1_tpl.py  原始.docx路径
"""
import docx, zipfile, shutil, re, sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

SRC = sys.argv[1] if len(sys.argv) > 1 else r'C:\Users\25534\Desktop\New folder\Duly Filled in Application Form.docx'
TPL = 'form1_tpl.docx'
shutil.copy(SRC, TPL)
d = docx.Document(TPL)
t = d.tables[0]

def uniq(row):
    out, seen = [], set()
    for c in row.cells:
        if id(c._tc) not in seen:
            seen.add(id(c._tc)); out.append(c)
    return out

def set_tok(cell, tok, append=False):
    """把单元格文本替换为 {{TOK}} 单 run，保留原段落样式。
    append=True 时在原文后追加一个 run。"""
    # 从单元格里找一个现成 run 的字体属性做克隆
    rpr_xml = None
    for p in cell.paragraphs:
        for r in p.runs:
            if r._r.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr') is not None:
                rpr_xml = r._r.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr')
                break
        if rpr_xml is not None:
            break
    p = cell.paragraphs[0]
    if append:
        run = p.add_run(' {{' + tok + '}}')
    else:
        # 清空所有段落、保留第一个段落及其 pPr
        for pp in cell.paragraphs[1:]:
            pp._p.getparent().remove(pp._p)
        for r in list(p.runs):
            p._p.remove(r._r)
        run = p.add_run('{{' + tok + '}}')
    if rpr_xml is not None:
        import copy
        run._r.insert(0, copy.deepcopy(rpr_xml))

R = lambda i: uniq(t.rows[i])
U = lambda row, idx: uniq(t.rows[row])[idx]

# ---- 勾选框（Symbol F0D6）：在对应空格放 {{CB_xxx}} ----
set_tok(U(4, 3), 'CB_EXPLORE'); set_tok(U(4, 5), 'CB_MINING')
set_tok(U(5, 3), 'CB_EXPL_LARGE'); set_tok(U(5, 5), 'CB_MINE_LARGE')
set_tok(U(6, 3), 'CB_EXPL_SMALL'); set_tok(U(6, 5), 'CB_MINE_SMALL')
set_tok(U(7, 3), 'CB_EXPL_SMALL2'); set_tok(U(7, 5), 'CB_ARTISANAL')
set_tok(U(10, 3), 'CB_INDIV'); set_tok(U(10, 5), 'CB_COOP')
set_tok(U(75, 3), 'CB_EXPL_ACT'); set_tok(U(76, 3), 'CB_MINE_ACT')

# ---- 单值字段 ----
set_tok(U(8, 2), 'MINERALS')
set_tok(U(9, 2), 'APPLICANT')
set_tok(U(11, 2), 'DOB')
set_tok(U(12, 2), 'NAT')
set_tok(U(13, 2), 'IDNO')
set_tok(U(14, 2), 'ADDR')
set_tok(U(15, 2), 'TEL')
set_tok(U(16, 2), 'MOBILE')
set_tok(U(17, 2), 'FAX')
set_tok(U(18, 2), 'EMAIL')
set_tok(U(19, 2), 'REGNO')

# ---- 股东（行21-25）/董事（27-31）：每行 idx2..5 = 名/股数/国籍/证件 ----
for i, r in enumerate(range(21, 26), 1):
    for j, s in enumerate(['N', 'S', 'NA', 'ID'], 2):
        try: set_tok(U(r, j), f'SH{i}_{s}')
        except IndexError: pass
for i, r in enumerate(range(27, 32), 1):
    for j, s in enumerate(['N', 'S', 'NA', 'ID'], 2):
        try: set_tok(U(r, j), f'DIR{i}_{s}')
        except IndexError: pass
set_tok(U(32, 3), 'CEO_N'); set_tok(U(32, 5), 'CEO_NA')
set_tok(U(33, 3), 'SEC_N'); set_tok(U(33, 5), 'SEC_NA')

# ---- 矿证历史表：现有(35-40)/曾持有(42-47)/子公司(49-54)/外国(56-61) ----
for i, r in enumerate(range(35, 41), 1):
    for j, s in enumerate(['NO', 'LOC', 'AR'], 2):
        try: set_tok(U(r, j), f'CUR{i}_{s}')
        except IndexError: pass
for i, r in enumerate(range(42, 48), 1):
    cs = uniq(t.rows[r])
    for j, s in enumerate(['NO', 'LOC', 'AR'], 2):
        try: set_tok(cs[j], f'PREV{i}_{s}')
        except IndexError: pass
for i, r in enumerate(range(49, 55), 1):
    for j, s in enumerate(['NO', 'LOC', 'AR'], 2):
        try: set_tok(U(r, j), f'SUB{i}_{s}')
        except IndexError: pass
for i, r in enumerate(range(56, 62), 1):
    for j, s in enumerate(['NO', 'LOC', 'AR'], 2):
        try: set_tok(U(r, j), f'FRN{i}_{s}')
        except IndexError: pass

# ---- Q13 曾申请记录（66-69）----
for i, r in enumerate(range(66, 70), 1):
    cs = uniq(t.rows[r])
    for j, s in enumerate(['LIC', 'LOC', 'AR', 'DT'], 1):
        try: set_tok(cs[j], f'HIST{i}_{s}')
        except IndexError: pass
# Q13 状态勾选：Grant/Rejection/Pending
for j, s in zip(range(2, 5), ['CB_GRANT', 'CB_REJECT', 'CB_PENDING']):
    try: set_tok(U(70, j), s, append=True)
    except IndexError: pass

# ---- Q14 地块 ----
set_tok(U(72, 3), 'PROVINCE'); set_tok(U(72, 5), 'DISTRICT')
set_tok(U(73, 3), 'LOCALITY')
set_tok(U(74, 2), 'SIZEKM2')
# 角点行 79-90：值格 = uniq idx4..15（idx16 是跨行合并续格，跳过），经纬 DMS 各占两格
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
def is_vmerge(c):
    return c._tc.find(f'{W}tcPr/{W}vMerge') is not None
for i, r in enumerate(range(79, 91), 1):
    cs = uniq(t.rows[r])
    vi = 0
    for j in range(4, len(cs)):
        if is_vmerge(cs[j]):
            continue
        vi += 1
        set_tok(cs[j], f'C{i}_V{vi}')

# ---- 费用 + 签名行 ----
set_tok(U(132, 1), 'FEE', append=True)

d.save(TPL)

# ---- XML 级：合并单元格内的 'None' 答案（Q12a/Q12b/Q13）+ 签名姓名 ----
zin = zipfile.ZipFile(TPL)
xml = zin.read('word/document.xml').decode('utf-8')
names = zin.namelist()
zin.close()
toks = ['Q12A', 'Q12B', 'Q13A']
for tok in toks:
    xml = xml.replace('>None<', '>{{' + tok + '}}<', 1)
# 签名行 "He XiaoJun" 被拼写检查拆成两个 run："He " + "XiaoJun"
xml = xml.replace('<w:t xml:space="preserve">He </w:t>', '<w:t>{{SIGN_NAME}} </w:t>', 1)
xml = xml.replace('<w:t>XiaoJun</w:t>', '<w:t></w:t>', 1)
zout = zipfile.ZipFile(TPL, 'w', zipfile.ZIP_DEFLATED)
zin2 = zipfile.ZipFile(SRC)
for n in zin2.namelist():
    zout.writestr(n, xml.encode('utf-8') if n == 'word/document.xml' else zin2.read(n))
zin2.close(); zout.close()

n_tok = len(re.findall(r'\{\{[A-Z0-9_]+\}\}', xml))
print('模板已生成:', TPL, '占位符数:', n_tok, '大小:', os.path.getsize(TPL))
